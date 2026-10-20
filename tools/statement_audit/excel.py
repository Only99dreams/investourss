"""Write the audit to an .xlsx workbook.

Three sheets, matching docs/financial-auditor-v2.md §5a:

* "Summary"        — the verified figures and how the statement was read.
* "Debit-Credit"   — one row per transaction the audit counted, with the amount
                     in a single Debit or Credit column and a TOTAL row. This
                     is exactly what was totalled, so the numbers are checkable
                     by hand.
* "As extracted"   — the raw cell table OpenCV/PDFium produced, so an
                     extraction problem is visible instead of hidden.
"""

from __future__ import annotations

from dataclasses import asdict, is_dataclass
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .audit_rules import FinancialAuditRulesResult

MONEY_FORMAT = "#,##0.00"

TITLE_FONT = Font(bold=True, size=14)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
LABEL_FONT = Font(bold=True)
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _jsonable(value: Any) -> Any:
    if is_dataclass(value):
        return {key: _jsonable(item) for key, item in asdict(value).items()}
    if isinstance(value, dict):
        return {key: _jsonable(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [_jsonable(item) for item in value]
    return value


def _autosize(worksheet, minimums: Optional[dict[int, int]] = None) -> None:
    minimums = minimums or {}
    for column in worksheet.columns:
        letter = get_column_letter(column[0].column)
        longest = max(
            (len(str(cell.value)) for cell in column if cell.value is not None), default=0
        )
        worksheet.column_dimensions[letter].width = max(minimums.get(column[0].column, 10), min(60, longest + 2))


def _write_summary_sheet(
    worksheet,
    audit: FinancialAuditRulesResult,
    reconciliation: dict,
    extraction: dict,
    printed_totals: dict,
    source: str,
    months_requested: float,
) -> None:
    worksheet["A1"] = "Financial audit"
    worksheet["A1"].font = TITLE_FONT
    worksheet["A2"] = source
    worksheet["A2"].font = Font(italic=True, color="595959")

    rows: list[tuple[str, Any]] = [
        ("Rules version", audit.rulesVersion),
        ("Period audited", f"{audit.periodStart} to {audit.periodEnd}"),
        ("Months audited", audit.auditMonths),
        ("Months requested", months_requested),
        ("Months in statement", audit.statementMonthsCovered),
        ("", ""),
        ("Total income", audit.totalIncome),
        ("Total expenses", audit.totalExpenses),
        ("Cash flow (not profit)", audit.cashFlow),
        ("Savings rate %", audit.savingsRate),
        ("Health score", audit.healthScore),
        ("Health status", audit.healthStatus),
        ("", ""),
        ("Transactions counted", len(audit.transactions)),
        ("Rows dropped (invalid)", audit.dataQuality.excludedInvalidTransactions),
        ("Duplicate rows ignored", audit.dataQuality.duplicateRowsIgnored),
        ("Undated transactions", audit.dataQuality.undatedTransactions),
        ("Undated excluded by window", audit.dataQuality.excludedUndatedTransactions),
        ("", ""),
        ("Opening balance", reconciliation.get("opening")),
        ("Closing balance", reconciliation.get("closing")),
        ("Credits (parser)", reconciliation.get("credits")),
        ("Debits (parser)", reconciliation.get("debits")),
        ("Expected closing", reconciliation.get("expectedClosing")),
        ("Difference", reconciliation.get("difference")),
        ("Reconciled", "yes" if reconciliation.get("reconciled") else "no"),
        ("", ""),
        ("Printed total credit", printed_totals.get("printedCredit")),
        ("Printed total debit", printed_totals.get("printedDebit")),
        ("Credit matches printed", printed_totals.get("creditMatches")),
        ("Debit matches printed", printed_totals.get("debitMatches")),
        ("", ""),
        ("Extraction method", extraction.get("method")),
        ("Pages", extraction.get("pages")),
        ("Table rows", extraction.get("rows")),
        ("Header row found", extraction.get("headerFound")),
        ("OCR pages", ", ".join(str(p) for p in extraction.get("ocrPages", [])) or "none"),
    ]
    start = 4
    for offset, (label, value) in enumerate(rows):
        worksheet.cell(row=start + offset, column=1, value=label).font = LABEL_FONT
        cell = worksheet.cell(row=start + offset, column=2, value=value)
        if label.startswith("Total") or label in ("Opening balance", "Closing balance", "Expected closing", "Difference", "Credits (parser)", "Debits (parser)"):
            cell.number_format = MONEY_FORMAT
        if label == "Savings rate %":
            cell.number_format = "0.00"

    warnings = list(extraction.get("warnings", []))
    if warnings:
        warning_start = start + len(rows) + 1
        worksheet.cell(row=warning_start, column=1, value="Notes").font = LABEL_FONT
        for offset, warning in enumerate(warnings):
            worksheet.cell(row=warning_start + 1 + offset, column=1, value=warning)

    _autosize(worksheet, {1: 26, 2: 28})


def _write_transactions_sheet(worksheet, audit: FinancialAuditRulesResult) -> None:
    worksheet["A1"] = "Debit / Credit (cleaned)"
    worksheet["A1"].font = TITLE_FONT
    worksheet["A2"] = (
        f"Period {audit.periodStart} to {audit.periodEnd} — "
        f"income {audit.totalIncome:,.2f}, expenses {audit.totalExpenses:,.2f} "
        f"({len(audit.transactions)} transaction rows counted)"
    )
    worksheet["A2"].font = Font(italic=True, color="595959")

    header = ["Date", "Description", "Category", "Debit", "Credit", "Reference"]
    header_row = 4
    for index, title in enumerate(header, start=1):
        cell = worksheet.cell(row=header_row, column=index, value=title)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = BORDER

    debits_total = 0.0
    credits_total = 0.0
    for offset, transaction in enumerate(audit.transactions):
        row = header_row + 1 + offset
        values = [
            transaction.date or "",
            transaction.description,
            transaction.category,
            transaction.amount if transaction.type == "debit" else None,
            transaction.amount if transaction.type == "credit" else None,
            transaction.transactionId or "",
        ]
        if transaction.type == "debit":
            debits_total += transaction.amount
        else:
            credits_total += transaction.amount
        for index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row, column=index, value=value)
            cell.border = BORDER
            if index in (4, 5):
                cell.number_format = MONEY_FORMAT

    total_row = header_row + 1 + len(audit.transactions)
    worksheet.cell(row=total_row, column=1, value="TOTAL").font = TOTAL_FONT
    worksheet.cell(row=total_row, column=2, value=f"{len(audit.transactions)} rows").font = TOTAL_FONT
    for column, value in ((4, debits_total), (5, credits_total)):
        cell = worksheet.cell(row=total_row, column=column, value=round(value, 2))
        cell.font = TOTAL_FONT
        cell.fill = TOTAL_FILL
        cell.number_format = MONEY_FORMAT
        cell.border = BORDER

    _autosize(worksheet, {1: 12, 2: 42, 3: 18, 4: 14, 5: 14, 6: 18})


def _write_raw_sheet(worksheet, raw_rows: list[list[Any]]) -> None:
    worksheet["A1"] = "As extracted"
    worksheet["A1"].font = TITLE_FONT
    worksheet["A2"] = "Exactly what OpenCV / PDFium read from the statement, before cleaning."
    worksheet["A2"].font = Font(italic=True, color="595959")
    for row_index, row in enumerate(raw_rows):
        for column_index, value in enumerate(row):
            worksheet.cell(row=row_index + 4, column=column_index + 1, value=value)
    _autosize(worksheet, {1: 14})


def write_workbook(
    path: str,
    *,
    source: str,
    audit: FinancialAuditRulesResult,
    reconciliation: dict,
    extraction: dict,
    printed_totals: dict,
    raw_rows: list[list[Any]],
    months_requested: float,
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    transactions = workbook.create_sheet("Debit-Credit")
    raw = workbook.create_sheet("As extracted")

    _write_summary_sheet(
        summary, audit, reconciliation, extraction, printed_totals, source, months_requested
    )
    _write_transactions_sheet(transactions, audit)
    _write_raw_sheet(raw, raw_rows)

    summary.sheet_view.showGridLines = False
    workbook.save(path)
